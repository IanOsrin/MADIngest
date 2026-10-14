"""scripts/mam-repeated-review-xlsx.py — the review list of repeated MAM songs, one row per COPY so the copies can be
compared side by side (Ian, 2026-10-14: "I think we need to redo the excel"). Read-only.

   python3 scripts/mam-repeated-review-xlsx.py ~/Downloads/mam-partial-dupes-3 tmp/vision-index.json out.xlsx
"""
import json, sys, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d, vix, out = sys.argv[1:4]
review = json.load(open(f'{d}/review.json')); B = json.load(open(f'{d}/backup-2026-10-14.json'))['records']
entries = {}
def collect(o):
    if isinstance(o, dict):
        p = o.get('path') or o.get('key') or o.get('rel')
        if isinstance(p, str) and o.get('size') is not None: entries[p] = o.get('size'); return
        for v in o.values(): collect(v)
    elif isinstance(o, list):
        for x in o:
            if isinstance(x, list) and len(x) >= 2 and isinstance(x[0], str): entries[x[0]] = x[1]
            else: collect(x)
collect(json.load(open(vix)))
def audio(u):
    if not u or not str(u).strip(): return ''
    p = '/gallo-' + u.split('/gallo-', 1)[1] if '/gallo-' in u else u
    s = entries.get(p)
    return f'yes ({s / 1e6:.0f} MB)' if s else 'link, file not found'
import unicodedata
asc = lambda x: unicodedata.normalize('NFKD', str(x or '')).encode('ascii', 'ignore').decode().lower()
art = lambda x: ' '.join(w for w in re.findall(r'[a-z0-9]+', re.sub(r'/\s*unknown|\bunknown\b', ' ', asc(x).replace('&', ' and '))) if w not in ('the', 'and'))
def secs(x):
    m = re.fullmatch(r'(?:(\d+):)?(\d{1,2}):(\d{2})', str(x or '').strip())
    return None if not m else int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + int(m.group(3))
def likely(rs):
    arts = {art(r.get('Track Artist')) for r in rs if art(r.get('Track Artist'))}
    if len(arts) > 1: return 'different songs — different artists'
    vers = {asc(r.get('Version')).strip() for r in rs if asc(r.get('Version')).strip()}
    if len(vers) > 1: return 'different versions of the song'
    ds = [secs(r.get('Duration')) for r in rs if secs(r.get('Duration'))]
    if len(ds) > 1 and max(ds) - min(ds) > 20: return 'different recordings — lengths differ by ' + str(max(ds) - min(ds)) + ' s'
    return 'check — could be the same recording twice'
kind = lambda r: r.split(':')[0].split(' — ')[0].split(' by ')[0].split(' (')[0]

wb = Workbook(); H = PatternFill('solid', fgColor='6D28D9'); KEEP = PatternFill('solid', fgColor='DCFCE7')
BANDS = [PatternFill('solid', fgColor='FFFFFF'), PatternFill('solid', fgColor='F3F4F6')]
top = Border(top=Side(style='medium', color='6D28D9'))
s = wb.active; s.title = 'Summary'
c = collections.Counter(kind(r) for g in review for r in g['reasons'])
for row in [['Repeated songs on MAM albums still to review — 14 Oct 2026 (after 868 duplicates were deleted)'], [],
            ['Songs to review', len(review)], ['Albums', len({g['cat'] for g in review})], ['Copies (records)', sum(g['copies'] for g in review)], [],
            ['Why they need a look (a song can have several reasons)']] + [['  ' + k, v] for k, v in c.most_common()]:
    s.append(row)
s['A1'].font = Font(bold=True, size=13); s.column_dimensions['A'].width = 80; s.column_dimensions['B'].width = 10

tally = collections.Counter()
ws = wb.create_sheet('Review')
COLS = ['Catalogue', 'Song', 'Likely', 'Why it needs a look', 'Copy', 'Record', 'Track #', 'Filename', 'ISRC', 'Duration', 'Version', 'Track Artist',
        'Audio on Vision', 'On website (MadStreamer)', 'Came from', 'Decision (keep / remove / both genuine)']
ws.append(COLS)
for c_ in ws[1]: c_.font = Font(bold=True, color='FFFFFF'); c_.fill = H; c_.alignment = Alignment(wrap_text=True, vertical='top')
for i, w in enumerate([14, 28, 30, 44, 6, 9, 7, 13, 14, 9, 14, 26, 18, 12, 10, 20], 1): ws.column_dimensions[get_column_letter(i)].width = w
for gi, g in enumerate(sorted(review, key=lambda g: (g['cat'], g['title'] or ''))):
    ids = [g['keep']] + g['drops']
    lk = likely([B[i] for i in ids]); tally[lk.split(' — ')[0].split(' of ')[0]] += 1
    for j, i in enumerate(ids):
        r = B[i]
        ws.append([g['cat'] if j == 0 else '', g['title'] if j == 0 else '', lk if j == 0 else '', ' · '.join(g['reasons']) if j == 0 else '',
                   'A (suggested keep)' if j == 0 else chr(66 + j - 1), int(i), r.get('Sequence Number') or '', r.get('Filename') or '',
                   r.get('ISRC') or '', r.get('Duration') or '', r.get('Version') or '', r.get('Track Artist') or '',
                   audio(r.get('Audio_Vision_URL')), 'yes' if r.get('Streamer_recid') else '', r.get('Sources') or '', ''])
        row = ws.max_row
        for col in range(1, len(COLS) + 1):
            ws.cell(row, col).fill = BANDS[gi % 2]
            ws.cell(row, col).alignment = Alignment(vertical='top', wrap_text=col in (3, 4))
            if j == 0: ws.cell(row, col).border = top
        ws.cell(row, 5).fill = KEEP if j == 0 else BANDS[gi % 2]
ws.freeze_panes = 'E2'; ws.auto_filter.ref = ws.dimensions
s.append([]); s.append(['Likely (from artist, version and length)'])
for k, v in tally.most_common(): s.append(['  ' + k, v])

rd = wb.create_sheet('Read me')
for l in ['Every song that still appears more than once on the same MAM album (same catalogue + same title), read live on 14 Oct 2026 — after today\'s 868 duplicate deletions. Placeholder titles ("unknown", "track 3"…) are left out: those are different songs.',
          'One row per copy; the copies of a song are grouped (alternating shading, a line above each song). Copy A is the suggested keeper (linked to the website, then ISRC, track number, audio).',
          '"Why it needs a look" says what differs between the copies: a different audio file, ISRC, duration, version, track number or artist suggests they may be genuinely different recordings — e.g. a reprise, a live take, or the same song placed twice on a compilation.',
          '"Likely" is a first guess from artist, version and length: different artists = different songs that share a title (keep both); different versions or lengths = different recordings; "check" = could really be the same recording twice.',
          'Fill the last column (keep / remove / both genuine) and send the file back: the "remove" ones will be checked and marked for deletion the usual way.']:
    rd.append([l])
rd.column_dimensions['A'].width = 150
for c_ in rd['A']: c_.alignment = Alignment(wrap_text=True)
wb.move_sheet('Review', offset=-1)
wb.active = 0
wb.save(out)
print(out, len(review), 'songs,', sum(g['copies'] for g in review), 'copies')
