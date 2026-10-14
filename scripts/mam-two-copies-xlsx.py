"""scripts/mam-two-copies-xlsx.py — for MAM albums whose songs are held twice and the second copy is NOT an empty
shell: what each copy holds, song by song, and whether one copy already contains everything (Ian, 2026-10-14:
"list what each copy holds for the 22 albums"). Read-only. Input: tmp/mam-two-copies.json (both copies read live).

   python3 scripts/mam-two-copies-xlsx.py tmp/mam-two-copies.json out.xlsx
"""
import json, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, out = sys.argv[1:3]
songs = json.load(open(src))
FIELDS = ['Sequence Number', 'Track Number', 'Filename', 'ISRC', 'Audio_Vision_URL', 'AudioHashSum', 'Duration', 'Version',
          'Track Artist', 'Composers', 'Composer', 'Publishers', 'cLine', 'pLine', 'Genre', 'Language', 'Catalogue_recid', 'CMS_recid', 'Streamer_recid']
SHORT = {'Sequence Number': 'Track #', 'Track Number': 'Track no.', 'Audio_Vision_URL': 'Audio on Vision', 'AudioHashSum': 'Audio hash',
         'Catalogue_recid': 'Gallo Catalogue link', 'CMS_recid': 'CMS 2024 link', 'Streamer_recid': 'MadStreamer link'}
JUNK = {'', '==', '0', 'none', 'null'}
has = lambda f, k: str((f or {}).get(k, '') if (f or {}).get(k) is not None else '').strip().lower() not in JUNK
val = lambda f, k: str((f or {}).get(k, '') if (f or {}).get(k) is not None else '').strip()
LINKS = {'Catalogue_recid', 'CMS_recid', 'Streamer_recid'}     # every copy links its own source record — not "data"

by = collections.OrderedDict()
for s in sorted(songs, key=lambda s: (s['cat'], s['title'] or '')): by.setdefault(s['cat'], []).append(s)

def verdict(ss):
    a_only = collections.Counter(); b_only = collections.Counter(); differ = collections.Counter()
    for s in ss:
        for k in FIELDS:
            if k in LINKS: continue
            ha, hb = has(s['a']['f'], k), has(s['b']['f'], k)
            if ha and not hb: a_only[k] += 1
            elif hb and not ha: b_only[k] += 1
            elif ha and hb and val(s['a']['f'], k) != val(s['b']['f'], k): differ[k] += 1
    if not b_only and not differ: v = 'Copy A has everything — copy B can go'
    elif not a_only and not differ: v = 'Copy B has everything — copy A can go'
    elif not differ: v = 'Each copy holds different parts — merge into one'
    else: v = 'Copies disagree on some values — check, then merge'
    return v, a_only, b_only, differ

fmt = lambda c: ', '.join(f'{SHORT.get(k, k)} ({n})' for k, n in c.most_common()) or '—'
wb = Workbook(); H = PatternFill('solid', fgColor='6D28D9'); A_ = PatternFill('solid', fgColor='DBEAFE'); B_ = PatternFill('solid', fgColor='FEF3C7'); X = PatternFill('solid', fgColor='FEE2E2')
s0 = wb.active; s0.title = 'Albums'
s0.append(['Catalogue', 'Songs', 'Copy A — record ids', 'Copy B — record ids', 'What to do', 'Only copy A has', 'Only copy B has', 'Both have, but different', 'Done ✓'])
for c in s0[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = H; c.alignment = Alignment(wrap_text=True, vertical='top')
for i, w in enumerate([16, 7, 22, 22, 40, 40, 40, 34, 8], 1): s0.column_dimensions[get_column_letter(i)].width = w
rng = lambda ids: (lambda x: f'{x[0]}–{x[-1]}' if x[-1] - x[0] == len(x) - 1 else ', '.join(map(str, x)))(sorted(int(i) for i in ids))
tally = collections.Counter()
for cat_, ss in by.items():
    v, ao, bo, df = verdict(ss); tally[v] += 1
    s0.append([cat_, len(ss), rng([s['a']['id'] for s in ss]), rng([s['b']['id'] for s in ss]), v, fmt(ao), fmt(bo), fmt(df), ''])
    for col in range(1, 10): s0.cell(s0.max_row, col).alignment = Alignment(wrap_text=True, vertical='top')
s0.freeze_panes = 'B2'; s0.auto_filter.ref = s0.dimensions

d = wb.create_sheet('Song by song')
cols = ['Catalogue', 'Song', 'Copy']
for k in FIELDS: cols.append(SHORT.get(k, k))
cols.insert(3, 'Record')
d.append(cols)
for c in d[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = H; c.alignment = Alignment(wrap_text=True, vertical='top')
for i, w in enumerate([14, 30, 6, 9] + [12] * len(FIELDS), 1): d.column_dimensions[get_column_letter(i)].width = w
for cat_, ss in by.items():
    for s in ss:
        for tag, cp, fill in (('A', s['a'], A_), ('B', s['b'], B_)):
            other = s['b'] if tag == 'A' else s['a']
            row = [cat_, s['title'], tag, int(cp['id'])] + [val(cp['f'], k)[:60] for k in FIELDS]
            d.append(row)
            d.cell(d.max_row, 3).fill = fill
            for j, k in enumerate(FIELDS, 5):
                if k in LINKS: continue
                mine, theirs = has(cp['f'], k), has(other['f'], k)
                if mine and not theirs: d.cell(d.max_row, j).fill = fill                     # only this copy has it
                elif mine and theirs and val(cp['f'], k) != val(other['f'], k): d.cell(d.max_row, j).fill = X
d.freeze_panes = 'E2'; d.auto_filter.ref = d.dimensions

rd = wb.create_sheet('Read me')
for l in [
    'MAM albums whose songs are held twice where the second copy is NOT an empty shell (the empty shells were deleted / marked separately). Both copies were read live on 14 Oct 2026. Nothing was changed.',
    'Copy A = the copy the duplicates list suggested keeping (most complete); copy B = the other. Record ids are MAM Songs records.',
    '"Albums": what each copy holds that the other lacks, with the number of songs. "What to do": if one copy holds everything the other does, the other can simply go; otherwise the copies need merging into one.',
    '"Song by song": every field for both copies. Blue = only copy A has it, yellow = only copy B has it, red = both have it but the values differ. Links to Gallo Catalogue / CMS 2024 / MadStreamer are shown but not counted — each copy links its own source record.',
    'Values equal to "==", "0" or blank count as empty.',
]: rd.append([l])
rd.column_dimensions['A'].width = 150
for c in rd['A']: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(out)
print(out, len(by), 'albums', dict(tally))
