"""scripts/isrc-cache-missing-songs.py — MadStreamer / MAM songs whose ALBUM is in the metadata cache but whose song
title isn't on it (Ian, 2026-10-13: "Album is in the cache but the song isn't list please"). Read-only.

For each song: the closest title on the cache album, and whether the song's current ISRC is on a cache song of that
album (then it is the same song under another spelling). Current values = the scans + today's fills and cache
changes (apply logs in ~/Downloads, versioned restores taken out).

   python3 scripts/isrc-cache-missing-songs.py tmp/isrc-gap-compare.json tmp/metadata-extract.json tmp/isrc-gap-mam.json out.xlsx
"""
import json, sys, os, re, unicodedata, collections, difflib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

cmp_, cache_f, mam_f, out = sys.argv[1:5]
def V(v):
    s = re.sub(r'[\s-]', '', str(v or '')).upper()
    return s if re.fullmatch(r'[A-Z]{2}[A-Z0-9]{3}\d{7}', s) else ''
cat = lambda c: re.sub(r'[^A-Za-z0-9]', '', str(c or '')).upper()
asc = lambda s: unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
strict = lambda s: re.sub(r'[^a-z0-9]+', '', asc(s))
loose = lambda s: re.sub(r'[^a-z0-9]+', '', re.sub(r'\(.*?\)|\[.*?\]', ' ', asc(s)))
words = lambda s: re.sub(r'[^a-z0-9]+', ' ', asc(s)).strip()

def overlay(*folders, drop=None):
    v = {}
    for f in folders:
        p = os.path.expanduser(f'~/Downloads/{f}/apply-log.jsonl')
        if not os.path.exists(p): continue
        for l in open(p):
            j = json.loads(l)
            if j.get('ok') and j.get('after'): v[j['recordId']] = j['after']
    for i in (drop or []): v.pop(i, None)
    return v
ms_drop = [r['id'] for r in json.load(open(os.path.expanduser('~/Downloads/madstreamer-cache-isrc/with-version.json')))] if os.path.exists(os.path.expanduser('~/Downloads/madstreamer-cache-isrc/with-version.json')) else []
ms_now = overlay('madstreamer-isrc-fill', 'madstreamer-cache-isrc', drop=ms_drop)
mam_now = overlay('mam-isrc-fill', 'mam-cache-isrc')

cache = json.load(open(cache_f)); cache = cache if isinstance(cache, list) else cache['rows']
alb = collections.defaultdict(list)
for c in cache:
    if c.get('catalogue') and c.get('track_name'): alb[cat(c['catalogue'])].append(c)

rows = [r for r in json.load(open(cmp_))['rows'] if r['cat'] or r['title']]
ms = [dict(id=str(r['msId']), cat=r['cat'], album=r['album'], title=r['title'], artist=r['artist'], seq=r['seq'], file=r['file'],
           cur=ms_now.get(str(r['msId'])) or r['ms'] or '', onSite=r['onSite']) for r in rows if r['where'] == 'MadStreamer']
mam = [dict(id=str(m['id']), cat=m['cat'], album='', title=m['title'], artist='', seq=m['seq'], file=m['file'],
            cur=mam_now.get(str(m['id'])) or V(m['isrcRaw']), onSite=False) for m in json.load(open(mam_f))]

def missing(songs):
    out_ = []
    for s in songs:
        a = alb.get(cat(s['cat']))
        if not a or not s['title']: continue
        if any(f(c['track_name']) and f(c['track_name']) == f(s['title']) for f in (strict, loose) for c in a): continue
        same = [c for c in a if s['cur'] and V(c.get('isrc')) == s['cur']]
        best = max(a, key=lambda c: difflib.SequenceMatcher(None, words(c['track_name']), words(s['title'])).ratio())
        sim = difflib.SequenceMatcher(None, words(best['track_name']), words(s['title'])).ratio()
        same_seq = [c for c in a if str(c.get('seq') or '').strip() and str(c.get('seq')).strip() == str(s['seq'] or '').strip()]
        if same: verdict, why = 'same song, other spelling', f"its ISRC is on \"{same[0]['track_name']}\" in the cache"
        elif sim >= 0.8: verdict, why = 'probably the same song, other spelling', f'{round(sim * 100)}% similar title'
        elif s['cur'] and any(V(c.get('isrc')) == s['cur'] for c in cache): verdict, why = 'ISRC is in the cache on ANOTHER album', ''
        else: verdict, why = 'not on the cache album', ''
        if verdict == 'ISRC is in the cache on ANOTHER album':
            o = next(c for c in cache if V(c.get('isrc')) == s['cur']); why = f"{o['track_name']} ({o['catalogue']})"
        out_.append(dict(s, cacheAlbum=a[0].get('album_title'), cacheSongs=len(a), best=best['track_name'], bestSeq=best.get('seq'), bestIsrc=best.get('isrc'),
                         sim=round(sim * 100), seqTitle=same_seq[0]['track_name'] if same_seq else '', verdict=verdict, why=why))
    return out_

res = {'MadStreamer': missing(ms), 'MAM': missing(mam)}

wb = Workbook(); H = PatternFill('solid', fgColor='6D28D9')
COLOR = {'same song, other spelling': 'DCFCE7', 'probably the same song, other spelling': 'FEF9C3', 'ISRC is in the cache on ANOTHER album': 'E0F2FE', 'not on the cache album': 'FEE2E2'}
ORDER = list(COLOR)
s = wb.active; s.title = 'Summary'
s.append(['Songs whose album IS in the metadata cache but the song is not (13 Oct 2026)']); s['A1'].font = Font(bold=True, size=13)
s.append([]); s.append(['', 'MadStreamer', 'MAM'])
for k in ORDER: s.append([k, sum(1 for x in res['MadStreamer'] if x['verdict'] == k), sum(1 for x in res['MAM'] if x['verdict'] == k)])
s.append(['Total', len(res['MadStreamer']), len(res['MAM'])])
s.append(['  …albums', len({cat(x['cat']) for x in res['MadStreamer']}), len({cat(x['cat']) for x in res['MAM']})])
s.append(['  …MadStreamer songs on the website now (ISRC + barcode, not hidden)', sum(1 for x in res['MadStreamer'] if x['onSite'] and x['cur']), ''])
s.append(['  …with no ISRC at all now', sum(1 for x in res['MadStreamer'] if not x['cur']), sum(1 for x in res['MAM'] if not x['cur'])])
for i, k in enumerate(ORDER, 4): s.cell(i, 1).fill = PatternFill('solid', fgColor=COLOR[k])
s.column_dimensions['A'].width = 75; s.column_dimensions['B'].width = 13; s.column_dimensions['C'].width = 13

COLS = ['Catalogue', 'Album', 'Song', 'Artist', '#', 'Filename', 'Record', 'ISRC now', 'On website?', 'What it looks like', 'Detail',
        'Closest title on the cache album', 'Its #', 'Its ISRC', 'Similarity %', 'Cache title at the same #', 'Cache album', 'Songs on cache album']
W = [14, 26, 30, 18, 4, 13, 9, 15, 8, 30, 36, 32, 5, 15, 9, 30, 26, 8]
for db in ('MadStreamer', 'MAM'):
    ws = wb.create_sheet(f'{db} ({len(res[db])})')
    ws.append(COLS)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = H; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(W, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'D2'
    for x in sorted(res[db], key=lambda x: (ORDER.index(x['verdict']), cat(x['cat']), str(x['seq']).zfill(4))):
        ws.append([x['cat'], x['album'] or x['cacheAlbum'], x['title'], x['artist'], x['seq'], x['file'], int(x['id']), x['cur'] or '—',
                   ('yes' if x['onSite'] and x['cur'] else 'no') if db == 'MadStreamer' else '', x['verdict'], x['why'],
                   x['best'], x['bestSeq'], x['bestIsrc'], x['sim'], x['seqTitle'], x['cacheAlbum'], x['cacheSongs']])
        ws.cell(ws.max_row, 10).fill = PatternFill('solid', fgColor=COLOR[x['verdict']])
    ws.auto_filter.ref = ws.dimensions

rd = wb.create_sheet('Read me')
for l in [
    'Songs in MadStreamer (website copy of ~03:00 today) and MAM (read today) whose catalogue IS in the metadata cache, but no song on that cache album has the same title (ignoring case, accents, punctuation, and bracketed bits). ISRCs shown are as they are now, after today\'s fills and cache changes. Nothing was changed.',
    '"same song, other spelling" — the song\'s ISRC is on a song of that cache album: only the title differs. Fix the title (in the cache or the database), no ISRC change needed.',
    '"probably the same song, other spelling" — no ISRC link, but the closest cache title is 80%+ similar. Check, then align the title.',
    '"ISRC is in the cache on ANOTHER album" — the ISRC belongs to a song on a different cache album (a compilation track, or a wrong ISRC).',
    '"not on the cache album" — the cache album simply doesn\'t list this song (bonus track, longer reissue, or a song on the wrong album). "Cache title at the same #" helps spot a song filed under the wrong album or a shifted track list.',
]: rd.append([l])
rd.column_dimensions['A'].width = 150
for c in rd['A']: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(out)
for db in res: print(db, len(res[db]), collections.Counter(x['verdict'] for x in res[db]).most_common())
print(out)
