"""scripts/mam-cache-gaps-xlsx.py — the workbook for scripts/mam-cache-gaps.mjs (MAM blanks the metadata cache can fill).
   python3 scripts/mam-cache-gaps-xlsx.py tmp/mam-cache-gaps.json ~/Downloads/MAM-gaps-fillable-from-cache.xlsx
"""
import json, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
d = json.load(open(src))
ALBUM_F = ['UPC', 'Label', 'Genre', 'Release Date', 'Year of Release', 'Album Title', 'Album Artist']
SONG_F = ['ISRC', 'Track Artist', 'Sequence Number', 'Duration', 'Genre', 'Composers', 'Producers', 'Publishers', 'Language',
          'pLine', 'cLine', 'Original Release Date', 'Rights Territories', 'Featured Artist', 'Lyrical Content Rating']
# Companion fields: MAM keeps a split / second copy of these (Track Number = Sequence Number, Composer 1-4 =
# Composers split into names, Producer = Producers, Language Code with Language). Most "gaps" are these —
# real blanks, but far less important than a missing ISRC or duration, so they get their own count.
COMPANION_F = ['Track Number', 'Composer', 'Composer 2', 'Composer 3', 'Composer 4', 'Producer', 'Language Code']

wb = Workbook()
def head(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='6D28D9')
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
def finish(ws, freeze): ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions

# Summary
s = wb.active; s.title = 'Summary'
ac = collections.Counter(k for a in d['albumGaps'] for k in a['fills'])
sc = collections.Counter(k for t in d['songGaps'] for k in t['fills'])
s.append(['MAM blanks the metadata cache can fill', '']); s['A1'].font = Font(bold=True, size=13)
for row in [
    ['MAM albums', d['mamAlbums']], ['MAM songs', d['mamSongs']], ['Metadata cache rows', d['cacheRows']],
    ['Albums with no rows in the cache (nothing to compare)', d['albumsWithoutCache']],
    ['MAM songs matched to a cache song by title', d['songsMatched']],
    ['Albums with at least one fillable blank', len(d['albumGaps'])],
    ['Songs with at least one fillable blank in a MAIN field', sum(1 for t in d['songGaps'] if any(f in t['fills'] for f in SONG_F))],
    ['Songs whose only fillable blanks are COMPANION fields (Composer 1-4, Track Number, Producer, Language Code)', sum(1 for t in d['songGaps'] if not any(f in t['fills'] for f in SONG_F))],
    ['Songs only NEARLY matching a cache title (check by hand)', len(d['nearMisses'])], [], ['ALBUM FIELD', 'Albums blank in MAM, filled in the cache']]:
    s.append(row)
for f in ALBUM_F: s.append([f, ac.get(f, 0)])
s.append([]); s.append(['SONG FIELD', 'Songs blank in MAM, filled in the cache'])
for f in SONG_F: s.append([f, sc.get(f, 0)])
s.append([]); s.append(['COMPANION FIELD', 'Songs blank in MAM, filled in the cache'])
for f in COMPANION_F: s.append([f, sc.get(f, 0)])
for r in s.iter_rows():
    if r[0].value in ('ALBUM FIELD', 'SONG FIELD', 'COMPANION FIELD'): r[0].font = r[1].font = Font(bold=True)
s.column_dimensions['A'].width = 58; s.column_dimensions['B'].width = 40

ws = wb.create_sheet('Album gaps')
head(ws, ['Catalogue', 'MAM Album', 'Album artist', 'MAM record', 'AlbumID', 'Blank fields'] + ALBUM_F, [16, 34, 24, 10, 10, 8] + [18] * len(ALBUM_F))
for a in sorted(d['albumGaps'], key=lambda a: a['cat']):
    ws.append([a['cat'], a['title'], a['artist'], int(a['recordId']), a['albumId'], len(a['fills'])] + [a['fills'].get(f, '') for f in ALBUM_F])
finish(ws, 'C2')

ws = wb.create_sheet('Song gaps')
head(ws, ['Catalogue', 'Album', 'MAM record', '#', 'Song', 'Main fields blank', 'Companion fields blank'] + SONG_F + COMPANION_F, [16, 30, 10, 5, 32, 9, 11] + [16] * (len(SONG_F) + len(COMPANION_F)))
# main-field gaps first, then companion-only songs
for t in sorted(d['songGaps'], key=lambda t: (not any(f in t['fills'] for f in SONG_F), t['cat'], str(t['seq'] or '').zfill(4))):
    ws.append([t['cat'], t['album'], int(t['recordId']), t['seq'], t['title'], sum(1 for f in SONG_F if f in t['fills']), sum(1 for f in COMPANION_F if f in t['fills'])]
              + [t['fills'].get(f, '') for f in SONG_F + COMPANION_F])
finish(ws, 'F2')

ws = wb.create_sheet('Near title matches')
head(ws, ['Catalogue', 'Album', 'MAM record', '#', 'MAM song title', 'Cache song title', 'Likeness', 'Would fill'] + SONG_F, [16, 28, 10, 5, 30, 30, 9, 9] + [16] * len(SONG_F))
for t in sorted(d['nearMisses'], key=lambda t: (t['cat'], str(t['seq'] or '').zfill(4))):
    ws.append([t['cat'], t['album'], int(t['recordId']), t['seq'], t['title'], t['cacheTitle'], t['score'], sum(1 for f in SONG_F if f in t['fills'])] + [t['fills'].get(f, '') for f in SONG_F])
finish(ws, 'F2')

rd = wb.create_sheet('Read me')
for line in [
    f"MAM (Music Arena Master) records with BLANK fields that the metadata cache (Ingrooves extract + ISRC register) has a value for. Read {d['at'][:16].replace('T', ' ')} UTC. Nothing was written.",
    'Same rules as the MAM tab\'s "Fill from cache": albums matched by catalogue number; songs paired by TITLE (ignoring case, accents, punctuation and bracketed bits), never by position; a field is listed only when it is EMPTY in MAM — values already in MAM are never compared or replaced here.',
    'Album gaps / Song gaps: one row per record; a value in a field column = MAM is blank there and that is the cache\'s value. Song gaps lists songs with a blank MAIN field first; after them come songs whose only blanks are COMPANION fields — MAM\'s split/second copies (Composer 1-4 from Composers, Track Number = Sequence Number, Producer = Producers, Language Code) — real but minor.',
    'Sequence numbers are only offered when the cache\'s numbers for that album look like a real track order (some cache albums carry running row numbers).',
    'Near title matches: the MAM title is close to a cache title but not the same (spelling drift, or a different song) — check by hand before using any of those values.',
    'To fill: the MAM tab → album → "Fill from cache" does one album at a time with a preview; a bulk fill can be built from this list on request.',
]: rd.append([line])
rd.column_dimensions['A'].width = 140
for c in rd['A']: c.alignment = Alignment(wrap_text=True)
wb.save(out)
print(out, 'albums', len(d['albumGaps']), 'songs', len(d['songGaps']), 'near', len(d['nearMisses']))
