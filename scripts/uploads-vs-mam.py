"""scripts/uploads-vs-mam.py — for every file the Gallo Master 01 → Vision copy put on Vision (tmp/gm01-copy-log.jsonl),
is the song in MAM, and is it linked to audio? (Ian, 2026-10-14: "identify which files are missing in MAM from the new
uploads".) Catalogue + title read from the Vision path <Artist>_<Album>_<CAT>/<Song>.wav; a side letter (BL 229A /
IAL 3008b) is tried without it too. Read-only.

   python3 scripts/uploads-vs-mam.py tmp/gm01-copy-log.jsonl tmp/mam-songs.json tmp/mam-albums.json out.xlsx
"""
import json, re, sys, os, unicodedata, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

logf, songsf, albumsf, out = sys.argv[1:5]
up = [json.loads(l) for l in open(logf) if l.strip()]
up = [l for l in up if l['result'] in ('copied', 'already there')]
S, A = json.load(open(songsf)), json.load(open(albumsf))
cat = lambda c: re.sub(r'[^A-Za-z0-9]', '', str(c or '')).upper()
asc = lambda s: unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
nt = lambda s: re.sub(r'[^a-z0-9]', '', asc(s))
loose = lambda s: re.sub(r'[^a-z0-9]', '', re.sub(r'\(.*?\)|\[.*?\]', ' ', asc(s)))
albums = {}
for a in A:
    for k in (cat(a.get('Album Catalogue Number')), cat(a.get('Reference Catalogue Number'))):
        if k and k not in albums: albums[k] = a
by_title = collections.defaultdict(list)          # MAM albums by album title — the same album under another catalogue
for a in A:
    if nt(a.get('Album Title')) and not re.match(r'(?i)singles?\b', str(a.get('Album Title'))): by_title[nt(a['Album Title'])].append(a)
songs = collections.defaultdict(list)
for s in S: songs[cat(s['Album Catalogue'])].append(s)
rows = []
for l in up:
    d = l['dest']; parts = d.split('/'); title = re.sub(r'^\d{1,2}[\s._-]+', '', os.path.splitext(parts[-1])[0]); folder = parts[-2]
    segs = [x for x in folder.split('_')[1:]][::-1] if '_' in folder else [re.sub(r'^.*?\s{2,}', '', folder)]
    raw = segs[0] if segs else ''
    keys = []
    for sg in segs:                                       # every catalogue-looking part, right to left
        k = cat(re.sub(r'\s*A\s*&\s*B$', '', sg, flags=re.I))
        keys += [k] + ([k[:-1]] if re.search(r'[0-9][AB]$', k) else [])
    keys = [k for k in keys if k]
    al = next((albums[k] for k in keys if k in albums), None)
    pool = []
    for k in set(keys + ([cat(al['Album Catalogue Number']), cat(al.get('Reference Catalogue Number'))] if al else [])): pool += songs.get(k, [])
    song = next((s for s in pool if nt(s['Track Name']) == nt(title)), None) or next((s for s in pool if loose(title) and loose(s['Track Name']) == loose(title)), None)
    if not al and not song: kind = '1 Album not in MAM'
    elif not song: kind = '2 Album in MAM, song not'
    elif not str(song.get('Audio_Vision_URL') or '').strip(): kind = '3 Song in MAM, no audio linked'
    elif d.split('/')[-1] in str(song.get('Audio_Vision_URL')): kind = '5 Song in MAM, linked to this file'
    else: kind = '4 Song in MAM, linked to other audio'
    ftitle = folder.split('_')[1] if folder.count('_') >= 2 else ''
    other = [a for a in by_title.get(nt(ftitle), []) if cat(a['Album Catalogue Number']) not in keys] if kind[0] in '12' and nt(ftitle) else []
    rows.append(dict(kind=kind, elsewhere='; '.join(f"{a['Album Catalogue Number']} ({a['Album Artist']})" for a in other[:3]), cat=raw, folder=folder, title=title, dest=d, mb=round(l['size'] / 1e6, 1), album=(al or {}).get('Album Title') or '',
                     albumCat=(al or {}).get('Album Catalogue Number') or '', albumId=(al or {}).get('id') or '', songId=(song or {}).get('id') or '',
                     songTitle=(song or {}).get('Track Name') or '', isrc=(song or {}).get('ISRC') or '', audio=(song or {}).get('Audio_Vision_URL') or ''))
wb = Workbook(); H = PatternFill('solid', fgColor='6D28D9')
COL = {'1': 'FEE2E2', '2': 'FFEDD5', '3': 'FEF9C3', '4': 'E0F2FE', '5': 'DCFCE7'}
s = wb.active; s.title = 'Summary'
s.append(['Files copied from Gallo Master 01 to Vision — are they in MAM? (' + str(len(rows)) + ' files so far)']); s['A1'].font = Font(bold=True, size=13); s.append([])
s.append(['', 'Files', 'Albums'])
for k, v in sorted(collections.Counter(r['kind'] for r in rows).items()):
    s.append([k[2:], v, len({r['cat'] for r in rows if r['kind'] == k})]); s.cell(s.max_row, 1).fill = PatternFill('solid', fgColor=COL[k[0]])
s.column_dimensions['A'].width = 48
for k in sorted({r['kind'] for r in rows}):
    ws = wb.create_sheet(k[2:30].replace(',', '')); rs = [r for r in rows if r['kind'] == k]
    cols = ['Catalogue (folder)', 'Folder on Vision', 'Song (file)', 'MB', 'Same album title in MAM under', 'MAM album', 'MAM catalogue', 'MAM album record', 'MAM song', 'MAM song record', 'ISRC', 'MAM audio link now', 'Vision file']
    ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = H
    for i, w in enumerate([14, 40, 30, 6, 30, 26, 13, 10, 28, 10, 14, 40, 70], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for r in sorted(rs, key=lambda r: (r['cat'], r['title'])):
        ws.append([r['cat'], r['folder'], r['title'], r['mb'], r['elsewhere'], r['album'], r['albumCat'], int(r['albumId']) if r['albumId'] else '', r['songTitle'], int(r['songId']) if r['songId'] else '', r['isrc'], r['audio'], r['dest']])
    ws.freeze_panes = 'D2'; ws.auto_filter.ref = ws.dimensions
rd = wb.create_sheet('Read me')
for t in ['Every file the Gallo Master 01 → Vision copy has put on Vision so far, checked against MAM (read 14 Oct). Catalogue and song title come from the Vision path …/<Artist>_<Album>_<Catalogue>/<Song>.wav; a side letter (A/B) is also tried without it. Nothing was changed.',
          '1 Album not in MAM — the catalogue is not in MAM. If "Same album title in MAM under" shows a catalogue, MAM probably has it as a reissue under that number; otherwise the album needs adding.',
          '2 Album in MAM, song not — MAM has the album but not this song (or under a different spelling).',
          '3 Song in MAM, no audio linked — the new Vision file can be linked to the song (MAM Audio_Vision_URL).',
          '4 Song in MAM, linked to other audio — MAM already points at another copy; this one is a second copy (often the same render elsewhere).',
          '5 Song in MAM, linked to this file — already done.']: rd.append([t])
rd.column_dimensions['A'].width = 140
wb.save(out)
json.dump(rows, open('tmp/upload-vs-mam.json', 'w'))
for k, v in sorted(collections.Counter(r['kind'] for r in rows).items()): print(f'  {k}: {v} files · {len({r["cat"] for r in rows if r["kind"] == k})} albums')
