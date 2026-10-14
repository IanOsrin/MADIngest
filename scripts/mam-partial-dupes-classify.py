"""scripts/mam-partial-dupes-classify.py — sort every repeated song on a MAM album into "clear-cut duplicate" (safe to
merge + mark the extra copy) or "needs review" (could be a genuine second recording), with the reason
(Ian, 2026-10-14: "run the check across all 536"). Read-only: writes the plan + a workbook; mam-partial-dupes.mjs applies.

A group = same catalogue + same title (ignoring case, accents, punctuation). It is CLEAR-CUT only if no two copies
  - carry different Version values,
  - point at different audio (Vision file of a different size),
  - carry different valid ISRCs,
  - differ in duration by more than 5 s,
  - both have a track number and the numbers differ,
  - are both linked to different MadStreamer records (the website would hold it twice too),
and every other value on the extra copy is either equal to the kept copy's or fills one of its empty fields.
Kept copy = linked to MadStreamer, then ISRC, track number, audio, then most fields filled, then lowest record id.

   python3 scripts/mam-partial-dupes-classify.py tmp/mam-partial-dupes.json tmp/vision-index.json ~/Downloads/mam-partial-dupes
"""
import json, sys, os, re, unicodedata, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, vix, outdir = sys.argv[1:4]
recs = json.load(open(src))['records']
nt = lambda s: re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower())
cat = lambda c: re.sub(r'[^A-Za-z0-9]', '', str(c or '')).upper()
JUNK = {'', '==', '0:00:00', '00:00:00', '0'}
blank = lambda v: v is None or str(v).strip() in JUNK
eq = lambda a, b: str(a if a is not None else '').strip() == str(b if b is not None else '').strip()
V = lambda v: (lambda s: s if re.fullmatch(r'[A-Z]{2}[A-Z0-9]{3}\d{7}', s) else '')(re.sub(r'[\s-]', '', str(v or '')).upper())
def secs(d):
    m = re.fullmatch(r'(?:(\d+):)?(\d{1,2}):(\d{2})', str(d or '').strip())
    return None if not m else int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + int(m.group(3))

# Vision sizes
entries = {}
def collect(o):
    if isinstance(o, dict):
        p = o.get('path') or o.get('key') or o.get('rel')
        if isinstance(p, str) and o.get('size') is not None: entries[p] = o.get('size'); return
        for v in o.values(): collect(v)
    elif isinstance(o, list):
        for x in o:
            if isinstance(x, list) and len(x) >= 2 and isinstance(x[0], str): entries[x[0]] = x[1]
            else: collect(x)
collect(json.load(open(vix)))
def vsize(u):
    if blank(u): return None
    p = '/gallo-' + u.split('/gallo-', 1)[1] if '/gallo-' in u else u
    return entries.get(p)

LINKS = {'id', 'MasterID', 'AlbumID', 'RecordingID', 'Sources', 'MatchMethod', 'CMS_recid', 'Catalogue_recid', 'Streamer_recid',
         'Timestamp', 'Conflicts', 'Albums::DDEX Status', 'Audio_Truth', 'Rights Territories', 'Track Number', 'Album Catalogue', 'Track Name'}
def score(r):
    return (bool(r.get('Streamer_recid')) * 16 + bool(V(r.get('ISRC'))) * 8 + (not blank(r.get('Sequence Number'))) * 4
            + bool(vsize(r.get('Audio_Vision_URL'))) * 2 + sum(1 for k, v in r.items() if k not in LINKS and not blank(v)) / 100, -int(r['id']))

groups = collections.defaultdict(list)
for r in recs:
    if nt(r.get('Track Name')): groups[(cat(r['Album Catalogue']), nt(r['Track Name']))].append(r)
groups = {k: v for k, v in groups.items() if len(v) > 1}
# placeholder titles are not songs: several DIFFERENT tracks called "unknown" are not duplicates
PLACEHOLDER = re.compile(r'^(unknown|untitled|notitle|tba|tbc|na|none|track\d*|song\d*|side[ab]\d*|\d+)$')
placeholders = {k: v for k, v in groups.items() if PLACEHOLDER.match(k[1])}
groups = {k: v for k, v in groups.items() if k not in placeholders}

plan, review = [], []
for (c, t), rs in sorted(groups.items()):
    reasons = []
    vers = {str(r.get('Version')).strip().lower() for r in rs if not blank(r.get('Version'))}
    if len(vers) > 1: reasons.append('different versions: ' + ' / '.join(sorted(vers)))
    paths = {str(r.get('Audio_Vision_URL')).strip() for r in rs if not blank(r.get('Audio_Vision_URL'))}
    sizes = [vsize(p_) for p_ in paths]
    if len(paths) > 1 and (None in sizes or len(set(sizes)) > 1):
        reasons.append('different audio files on Vision' if None not in sizes else 'different audio paths, not all found on Vision — same audio not proven')
    files = {str(r.get('Filename')).strip().upper() for r in rs if not blank(r.get('Filename'))}
    if len(files) > 1 and not (len(paths) >= 1 and len(set(sizes)) == 1 and None not in sizes and all(not blank(r.get('Audio_Vision_URL')) for r in rs if not blank(r.get('Filename')))):
        reasons.append('different asset numbers (' + ', '.join(sorted(files)) + ') and the audio cannot be compared')
    isrcs = {V(r.get('ISRC')) for r in rs} - {''}
    if len(isrcs) > 1: reasons.append('different ISRCs: ' + ', '.join(sorted(isrcs)))
    durs = [secs(r.get('Duration')) for r in rs if secs(r.get('Duration'))]
    if durs and max(durs) - min(durs) > 5: reasons.append(f'durations differ by {max(durs) - min(durs)} s')
    seqs = {str(r.get('Sequence Number')).strip() for r in rs if not blank(r.get('Sequence Number'))}
    if len(seqs) > 1: reasons.append('placed at different track numbers: ' + ', '.join(sorted(seqs, key=lambda x: (len(x), x))))
    ms = {str(r.get('Streamer_recid')) for r in rs if not blank(r.get('Streamer_recid'))}
    if len(ms) > 1: reasons.append('each copy linked to its own MadStreamer record (on the website twice)')
    keep = max(rs, key=score); drops = [r for r in rs if r is not keep]
    sets, clash = {}, []
    if not reasons:
        for d in drops:
            for k, v in d.items():
                if k in LINKS or blank(v): continue
                if blank(keep.get(k)) and blank(sets.get(k)): sets[k] = v
                elif not eq(v, sets.get(k, keep.get(k))) and k not in ('Filename', 'Audio_Vision_URL', 'Audio_S3_URL', 'CMS_Audio_File', 'AudioHashSum'):
                    clash.append(f'{k}: "{str(keep.get(k))[:30]}" vs "{str(v)[:30]}"')
            # a different filename / audio path is fine only when the audio is the same (checked above) or the drop's is absent
        if clash: reasons.append('values disagree — ' + '; '.join(sorted(set(clash))[:4]))
    row = dict(cat=rs[0]['Album Catalogue'], title=keep['Track Name'], copies=len(rs), keep=str(keep['id']), drops=[str(d['id']) for d in drops],
               set={k: v for k, v in sets.items()}, expect={k: keep.get(k, '') for k in sets}, reasons=reasons)
    (review if reasons else plan).append(row)

os.makedirs(outdir, exist_ok=True)
json.dump({'records': {str(r['id']): r for r in recs}}, open(f'{outdir}/backup-2026-10-14.json', 'w'))
json.dump(plan, open(f'{outdir}/plan.json', 'w'), indent=1, ensure_ascii=False)

wb = Workbook(); H = PatternFill('solid', fgColor='6D28D9')
s = wb.active; s.title = 'Summary'
why = collections.Counter(r.split(':')[0].split(' — ')[0].split(' by ')[0].split(' (')[0] for g in review for r in g['reasons'])
for row in [['MAM albums where some songs are held more than once — 14 Oct 2026'], [],
            ['Albums checked', len({cat(g['cat']) for g in plan + review})], ['Repeated songs (groups)', len(plan) + len(review)],
            [], ['CLEAR-CUT duplicates (extra copy marked for deletion)', len(plan)], ['  …extra records', sum(len(g['drops']) for g in plan)],
            ['  …kept copies that receive values from the extra copy', sum(1 for g in plan if g['set'])],
            ['NEEDS REVIEW (could be a genuine second recording)', len(review)],
            ['Placeholder titles ("unknown" etc.) — different songs, NOT duplicates, left alone', f'{len(placeholders)} titles / {sum(len(v) for v in placeholders.values())} records'], []] + [['  ' + k, v] for k, v in why.most_common()]:
    s.append(row)
s['A1'].font = Font(bold=True, size=13); s.column_dimensions['A'].width = 70; s.column_dimensions['B'].width = 10
for name, rows, cols in (('Clear-cut', plan, ['Catalogue', 'Song', 'Copies', 'KEEP record', 'REMOVE record(s)', 'Copied onto kept copy']),
                         ('Needs review', review, ['Catalogue', 'Song', 'Copies', 'Suggested keep', 'Other record(s)', 'Why it needs a look'])):
    ws = wb.create_sheet(name); ws.append(cols)
    for c_ in ws[1]: c_.font = Font(bold=True, color='FFFFFF'); c_.fill = H
    for i, w in enumerate([16, 34, 7, 12, 22, 90], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for g in rows:
        ws.append([g['cat'], g['title'], g['copies'], int(g['keep']), ', '.join(g['drops']),
                   ', '.join(f'{k} = {str(v)[:40]}' for k, v in g['set'].items()) if name == 'Clear-cut' else ' · '.join(g['reasons'])])
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
rd = wb.create_sheet('Read me')
for l in ['Every song that appears more than once on the same MAM album (same catalogue + same title), read live 14 Oct 2026.',
          'CLEAR-CUT: the copies cannot be two different recordings — no different version, audio file, ISRC, duration (>5 s) or track number — and the extra copy holds nothing that contradicts the kept one. Anything it holds that the kept copy lacks is copied across first; then the extra copy is marked "DELETE: empty duplicate 2026-10-14 (8)".',
          'NEEDS REVIEW: something suggests the copies may be genuinely different (a reprise, live take, remaster, or the same song placed twice on a compilation), or their details disagree. Nothing was changed for these.']:
    rd.append([l])
rd.column_dimensions['A'].width = 150
for c_ in rd['A']: c_.alignment = Alignment(wrap_text=True)
wb.save(os.path.expanduser('~/Downloads/MAM-repeated-songs-check-2026-10-14.xlsx'))
print('placeholder titles skipped', len(placeholders), '| groups', len(plan) + len(review), '| clear-cut', len(plan), 'extra records', sum(len(g['drops']) for g in plan), '| review', len(review))
print(why.most_common())
