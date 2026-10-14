"""scripts/cms-double-check-report.py — report on the CMS 2024 songs marked Statusx = "Double check": their state in CMS,
MAM and MadStreamer, and whether they are on the website (Ian, 2026-10-14). Read-only.

Matching — MAM: the MAM record whose CMS_recid is this CMS record, else same catalogue + Filename, else catalogue +
title. MadStreamer (nightly copy): same ISRC (preferring the same catalogue), else catalogue + Filename, else
catalogue + title. "On the website" = MadStreamer has it with a valid ISRC, a barcode, a cover and not hidden.

   python3 scripts/cms-double-check-report.py tmp/cms-double-check.json tmp/metadata-extract.json tmp/vision-index.json out.xlsx
"""
import json, sys, re, unicodedata, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, cache_f, vix, out = sys.argv[1:5]
D = json.load(open(src)); cms, mam, ms = D['cms'], D['mam'], D['ms']
nt = lambda s: re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower())
cat = lambda c: re.sub(r'[^A-Za-z0-9]', '', str(c or '')).upper()
fk = lambda f: re.sub(r'\.(wav|mp3)$', '', str(f or '').strip(), flags=re.I).upper()
V = lambda v: (lambda s: s if re.fullmatch(r'[A-Z]{2}[A-Z0-9]{3}\d{7}', s) else '')(re.sub(r'[\s-]', '', str(v or '')).upper())

cache = json.load(open(cache_f)); cache = cache if isinstance(cache, list) else cache['rows']
in_cache = {V(r.get('isrc')) for r in cache if V(r.get('isrc'))}
entries = {}
def collect(o):
    if isinstance(o, dict):
        p = o.get('path') or o.get('key') or o.get('rel')
        if isinstance(p, str) and o.get('size') is not None: entries[p] = o.get('size'); return
        for v in o.values(): collect(v)
    elif isinstance(o, list):
        for x in o:
            if isinstance(x, list) and len(x) >= 2 and isinstance(x[0], str): entries[x[0]] = x[1]
            else: collect(x)
collect(json.load(open(vix)))
on_vision = lambda u: bool(u) and ('/gallo-' + u.split('/gallo-', 1)[1] if '/gallo-' in u else u) in entries

mam_by_cms = collections.defaultdict(list); mam_by_file = collections.defaultdict(list); mam_by_title = collections.defaultdict(list)
for m in mam:
    if m.get('CMS_recid'): mam_by_cms[str(m['CMS_recid'])].append(m)
    if m.get('Filename'): mam_by_file[(cat(m['Album Catalogue']), fk(m['Filename']))].append(m)
    mam_by_title[(cat(m['Album Catalogue']), nt(m['Track Name']))].append(m)
ms_by_isrc = collections.defaultdict(list); ms_by_file = collections.defaultdict(list); ms_by_title = collections.defaultdict(list)
for t in ms:
    if V(t['isrc']): ms_by_isrc[V(t['isrc'])].append(t)
    if t['file']: ms_by_file[(cat(t['cat']), fk(t['file']))].append(t)
    ms_by_title[(cat(t['cat']), nt(t['title']))].append(t)

import os
_rem = os.path.expanduser('~/Downloads/madstreamer-upc-removal/backup-2026-10-13.json')
upc_removed = {r['recordId'] for r in json.load(open(_rem))['records'] if str(r.get('upc') or '').strip()} if os.path.exists(_rem) else set()
rows = []
for c in cms:
    C = cat(c.get('Album Catalogue Number') or c.get('Reference Catalogue Number')); isrc = V(c.get('ISRC'))
    # MAM
    m, how = None, ''
    for key, idx, label in ((str(c['id']), mam_by_cms, 'CMS link'), ((C, fk(c.get('Filename'))), mam_by_file, 'catalogue + filename'), ((C, nt(c.get('Track Name'))), mam_by_title, 'catalogue + title')):
        if idx.get(key): m, how = idx[key][0], label; break
    mam_copies = len(mam_by_title.get((C, nt(c.get('Track Name'))), [])) if m else 0
    if not m: mam_status = 'NOT in MAM'
    elif not V(m.get('ISRC')): mam_status = 'in MAM — no ISRC'
    elif V(m.get('ISRC')) != isrc: mam_status = f'in MAM — DIFFERENT ISRC ({V(m.get("ISRC"))})'
    else: mam_status = 'in MAM — same ISRC'
    mam_audio = 'yes' if m and on_vision(m.get('Audio_Vision_URL')) else ('link, file not found' if m and m.get('Audio_Vision_URL') else ('—' if m else ''))
    # MadStreamer
    cands = ms_by_isrc.get(isrc, [])
    t = next((x for x in cands if cat(x['cat']) == C), None) or (cands[0] if cands else None)
    mhow = 'ISRC' if t else ''
    if not t:
        for key, idx, label in (((C, fk(c.get('Filename'))), ms_by_file, 'catalogue + filename'), ((C, nt(c.get('Track Name'))), ms_by_title, 'catalogue + title')):
            if idx.get(key): t, mhow = idx[key][0], label; break
    if not t: ms_status, site, why = 'NOT in MadStreamer', 'no', 'not in MadStreamer'
    else:
        miss = []
        if not V(t['isrc']): miss.append('no ISRC')
        elif V(t['isrc']) != isrc: miss.append(f'different ISRC {V(t["isrc"])}')
        if not t['upc']: miss.append('no barcode (removed 13 Oct: album had partial ISRCs)' if str(t['id']) in upc_removed else 'no barcode')
        if not t['has_art']: miss.append('no cover')
        if (t['visibility'] or '').lower() == 'hide': miss.append('hidden')
        if not t['has_audio']: miss.append('no audio')
        site = 'yes' if not [x for x in miss if not x.startswith('different ISRC') and x != 'no audio'] and V(t['isrc']) else 'no'
        ms_status = 'in MadStreamer' + (f' as {t["cat"]}' if cat(t['cat']) != C else '')
        why = ', '.join(miss) or 'all good'
    cms_audio = str(c.get('Audio Test') or '').strip()
    cms_audio = 'BadAudio' if cms_audio == 'BadAudio' else ('Valid' if cms_audio.startswith('Valid') else ('tested (format details only)' if cms_audio.startswith('Format') else (cms_audio[:20] or '—')))
    rows.append(dict(cat=c.get('Album Catalogue Number') or c.get('Reference Catalogue Number'), album=c.get('Album Title'), artist=c.get('Track Artist'),
        seq=c.get('Sequence Number'), title=c.get('Track Name'), isrc=isrc, upc=c.get('UPC'), file=c.get('Filename'), cmsId=c['id'],
        cmsVis=c.get('Visibility') or '', cmsAudio=cms_audio, cmsErr=str(c.get('error') or ''), inCache='yes' if isrc in in_cache else 'no',
        mamId=m['id'] if m else '', mamHow=how, mamStatus=mam_status, mamAudio=mam_audio, mamCopies=mam_copies,
        msId=t['id'] if t else '', msHow=mhow, msStatus=ms_status, site=site, why=why))

def verdict(r):
    if r['cmsAudio'] == 'BadAudio': return '1 Bad audio in CMS'
    if r['mamStatus'].startswith('NOT') and r['msStatus'].startswith('NOT'): return '2 In CMS only'
    if 'DIFFERENT ISRC' in r['mamStatus'] or 'different ISRC' in r['why']: return '3 ISRC disagrees'
    if r['site'] == 'yes': return '6 On the website'
    if r['msStatus'].startswith('NOT'): return '4 Not in MadStreamer'
    return '5 In MadStreamer, not shown'
for r in rows: r['verdict'] = verdict(r)

wb = Workbook(); H = PatternFill('solid', fgColor='6D28D9')
COL = {'1': 'FEE2E2', '2': 'FFEDD5', '3': 'FEF3C7', '4': 'FEF9C3', '5': 'E0F2FE', '6': 'DCFCE7'}
s = wb.active; s.title = 'Summary'
s.append(['CMS 2024 songs marked Statusx = "Double check" — status in CMS, MAM and MadStreamer (14 Oct 2026)']); s['A1'].font = Font(bold=True, size=13)
s.append([]); s.append(['Songs', len(rows)]); s.append(['Albums', len({cat(r['cat']) for r in rows})])
s.append(['All from delivery batch', 'DDEX Asset2024-02-16']); s.append([])
s.append(['Overall', 'Songs', 'Albums'])
vc = collections.Counter(r['verdict'] for r in rows)
for k in sorted(vc):
    s.append([k[2:], vc[k], len({cat(r['cat']) for r in rows if r['verdict'] == k})]); s.cell(s.max_row, 1).fill = PatternFill('solid', fgColor=COL[k[0]])
s.append([])
for title, key in (('CMS audio test', 'cmsAudio'), ('MAM', 'mamStatus'), ('MadStreamer', 'msStatus'), ('On the website', 'site'), ('Why not on the website', 'why'), ('ISRC in the metadata cache (Ingrooves)', 'inCache')):
    s.append([title]); s.cell(s.max_row, 1).font = Font(bold=True)
    cc = collections.Counter(re.sub(r' \(.*\)$| as .*$|ZA[A-Z0-9]+|US[A-Z0-9]+', '', str(r[key])) if key in ('mamStatus', 'msStatus') else r[key] for r in rows) if key != 'why' else \
         collections.Counter(x.strip() for r in rows if r['site'] == 'no' and not r['msStatus'].startswith('NOT') for x in re.sub(r'different ISRC [A-Z0-9]+', 'different ISRC', r['why'].replace(': album had partial ISRCs', '')).split(','))
    for k, v in cc.most_common(): s.append(['  ' + str(k), v])
    s.append([])
s.column_dimensions['A'].width = 60; s.column_dimensions['B'].width = 10; s.column_dimensions['C'].width = 10

a = wb.create_sheet('Albums')
a.append(['Catalogue', 'Album', 'Artist', 'Songs', 'Bad audio (CMS)', 'In MAM', 'In MadStreamer', 'On the website', 'Main issue'])
for c_ in a[1]: c_.font = Font(bold=True, color='FFFFFF'); c_.fill = H
for i, w in enumerate([15, 32, 26, 7, 10, 8, 12, 12, 40], 1): a.column_dimensions[get_column_letter(i)].width = w
by = collections.defaultdict(list)
for r in rows: by[cat(r['cat'])].append(r)
for k, rs in sorted(by.items(), key=lambda kv: (min(r['verdict'] for r in kv[1]), kv[0])):
    main = collections.Counter(r['verdict'][2:] for r in rs).most_common(1)[0][0]
    a.append([rs[0]['cat'], rs[0]['album'], rs[0]['artist'], len(rs), sum(r['cmsAudio'] == 'BadAudio' for r in rs), sum(not r['mamStatus'].startswith('NOT') for r in rs),
              sum(not r['msStatus'].startswith('NOT') for r in rs), sum(r['site'] == 'yes' for r in rs), main])
    a.cell(a.max_row, 9).fill = PatternFill('solid', fgColor=COL[min(r['verdict'] for r in rs)[0]])
a.freeze_panes = 'C2'; a.auto_filter.ref = a.dimensions

w = wb.create_sheet('Songs')
COLS = ['Overall', 'Catalogue', 'Album', '#', 'Song', 'Artist', 'ISRC', 'Barcode', 'Filename', 'CMS record', 'CMS visibility', 'CMS audio test',
        'MAM record', 'MAM status', 'MAM matched by', 'MAM audio on Vision', 'MAM copies of this song', 'MadStreamer record', 'MadStreamer status',
        'On the website', 'Why not / notes', 'ISRC in metadata cache']
w.append(COLS)
for c_ in w[1]: c_.font = Font(bold=True, color='FFFFFF'); c_.fill = H; c_.alignment = Alignment(wrap_text=True, vertical='top')
for i, wd in enumerate([20, 14, 26, 4, 28, 22, 14, 14, 12, 9, 9, 10, 9, 26, 16, 12, 9, 10, 22, 8, 34, 9], 1): w.column_dimensions[get_column_letter(i)].width = wd
for r in sorted(rows, key=lambda r: (r['verdict'], cat(r['cat']), int(re.sub(r'\D', '', str(r['seq'])) or 0))):
    w.append([r['verdict'][2:], r['cat'], r['album'], r['seq'], r['title'], r['artist'], r['isrc'], r['upc'], r['file'], int(r['cmsId']), r['cmsVis'], r['cmsAudio'],
              int(r['mamId']) if r['mamId'] else '', r['mamStatus'], r['mamHow'], r['mamAudio'], r['mamCopies'] or '', int(r['msId']) if r['msId'] else '', r['msStatus'],
              r['site'], r['why'], r['inCache']])
    w.cell(w.max_row, 1).fill = PatternFill('solid', fgColor=COL[r['verdict'][0]])
w.freeze_panes = 'F2'; w.auto_filter.ref = w.dimensions

rd = wb.create_sheet('Read me')
for l in ['The 1,801 songs in Gallo CMS 2024 (layout Song Files) whose Statusx field is "Double check" — all from delivery batch DDEX Asset2024-02-16, all with an ISRC and a barcode in CMS. Read 14 Oct 2026; nothing was changed.',
          'MAM: matched through its CMS_recid link to the CMS record, else same catalogue + Filename, else catalogue + title. "MAM copies" > 1 = MAM holds the song more than once.',
          'MadStreamer (the website\'s database, nightly copy of ~03:00 today): matched by ISRC (preferring the same catalogue — it may sit under a variant such as "BL 156A"), else catalogue + Filename, else catalogue + title.',
          'On the website = MadStreamer has it with a valid ISRC, a barcode, a cover and not hidden (the rule the site uses). "Why not" lists what is missing.',
          'Overall, in order of urgency: Bad audio in CMS (CMS\'s own Audio Test field) · In CMS only (not in MAM or MadStreamer) · ISRC disagrees between the databases · Not in MadStreamer · In MadStreamer but not shown · On the website.',
          '"no barcode (removed 13 Oct…)" = we cleared the barcode on purpose on 13 Oct because the album had songs without ISRCs (backup in ~/Downloads/madstreamer-upc-removal); it can go back once the album is complete.',
          '"ISRC in metadata cache" = the Ingrooves extract / ISRC register holds that ISRC.']:
    rd.append([l])
rd.column_dimensions['A'].width = 150
for c_ in rd['A']: c_.alignment = Alignment(wrap_text=True)
wb.save(out)
print(out); print(sorted(vc.items()))
print('MAM:', collections.Counter(re.sub(r' \(.*\)$', '', r['mamStatus']) for r in rows).most_common())
print('MS:', collections.Counter(re.sub(r' as .*$', '', r['msStatus']) for r in rows).most_common(), '| on site', sum(r['site'] == 'yes' for r in rows))
