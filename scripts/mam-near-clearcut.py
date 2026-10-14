"""scripts/mam-near-clearcut.py — second batch of repeated MAM songs (Ian, 2026-10-14: "do the near-clear-cut ones as a
second batch"): songs from the review list whose ONLY problem was "values disagree", accepted only when EVERY
disagreement is cosmetic (same artist written differently, © / ℗ symbols, composer credits with or without roles,
0:05:25 vs 00:05:25, "zu" vs "Zulu", date formats …). The better-formatted value is kept. Read-only: writes plan.json.

   python3 scripts/mam-near-clearcut.py ~/Downloads/MAM-repeated-songs-check-2026-10-14.xlsx ~/Downloads/mam-partial-dupes/backup-2026-10-14.json ~/Downloads/mam-partial-dupes-2
"""
import json, sys, re, unicodedata, collections, datetime
from openpyxl import load_workbook

xlsx, backup, outdir = sys.argv[1:4]
B = json.load(open(backup))['records']
JUNK = {'', '==', '0:00:00', '00:00:00', '0'}
blank = lambda v: v is None or str(v).strip() in JUNK
asc = lambda s: unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
NUM = {'one': '1', 'two': '2', 'three': '3', 'four': '4', 'five': '5', 'six': '6', 'seven': '7', 'eight': '8', 'nine': '9', 'ten': '10'}
def artist(s):
    s = asc(s).replace('&', ' and ')
    s = re.sub(r'/\s*unknown\b|\bunknown\b', ' ', s)
    w = [NUM.get(x, x) for x in re.findall(r'[a-z0-9]+', s) if x not in ('the', 'and')]
    return ' '.join(w)
def secs(d):
    m = re.fullmatch(r'(?:(\d+):)?(\d{1,2}):(\d{2})', str(d or '').strip())
    return None if not m else int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + int(m.group(3))
def date(s):
    s = str(s or '').strip()
    for f in ('%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y'):
        try: return datetime.datetime.strptime(s, f).date()
        except ValueError: pass
    return s
line = lambda s: re.sub(r'[^a-z0-9]', '', asc(s).replace('(p)', '').replace('(c)', ''))
names = lambda s: sorted(set(re.sub(r'\s+', ' ', x).strip() for x in re.sub(r'<[^>]*>', '', asc(s)).split(',') if x.strip()))
RATING = {'notexplicit': 'clean', 'no': 'clean', 'clean': 'clean', 'n': 'clean', 'explicit': 'explicit', 'yes': 'explicit', 'y': 'explicit'}
LANG = {'zu': 'zulu', 'xh': 'xhosa', 'af': 'afrikaans', 'en': 'english', 'st': 'sotho', 'tn': 'tswana', 'ts': 'tsonga', 've': 'venda', 'ss': 'swati', 'nr': 'ndebele', 'nso': 'pedi'}
lang = lambda s: LANG.get(asc(s).strip(), asc(s).strip())
SAME = {
    'Track Artist': lambda a, b: artist(a) == artist(b),
    'Genre': lambda a, b: re.sub(r'[^a-z]', '', asc(a)) == re.sub(r'[^a-z]', '', asc(b)),
    'Local Genre': lambda a, b: re.sub(r'[^a-z]', '', asc(a)) == re.sub(r'[^a-z]', '', asc(b)),
    'Duration': lambda a, b: secs(a) is not None and abs(secs(a) - secs(b or '')) <= 1 if secs(b) is not None else False,
    'pLine': lambda a, b: line(a) == line(b), 'cLine': lambda a, b: line(a) == line(b),
    'Lyrical Content Rating': lambda a, b: RATING.get(asc(a).replace(' ', '')) == RATING.get(asc(b).replace(' ', '')) is not None,
    'Original Release Date': lambda a, b: date(a) == date(b),
    'Language': lambda a, b: lang(a) == lang(b), 'Language Code': lambda a, b: lang(a) == lang(b),
    'Composers': lambda a, b: names(a) == names(b), 'Composer': lambda a, b: names(a) == names(b),
    'Publishers': lambda a, b: names(a) == names(b),
}
def better(k, a, b):          # which of two cosmetically-equal values to keep
    if k == 'Track Artist': return b if re.search(r'unknown', str(a), re.I) and not re.search(r'unknown', str(b), re.I) else a
    if k == 'pLine': return a if str(a).strip().startswith('℗') or not str(b).strip().startswith('℗') else b
    if k == 'cLine': return a if str(a).strip().startswith('©') or not str(b).strip().startswith('©') else b
    if k in ('Composers', 'Composer', 'Publishers'): return a if len(str(a)) >= len(str(b)) else b
    if k == 'Duration': return a if re.fullmatch(r'\d{2}:\d{2}:\d{2}', str(a).strip()) or not re.fullmatch(r'\d{2}:\d{2}:\d{2}', str(b).strip()) else b
    if k == 'Language Code': return a if len(str(a).strip()) <= 3 else b
    if k == 'Lyrical Content Rating': return a if str(a).strip()[:1].isupper() and len(str(a)) > 2 else b
    if k == 'Original Release Date': return a if re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(a).strip()) else b
    return a
LINKS = {'id', 'MasterID', 'AlbumID', 'RecordingID', 'Sources', 'MatchMethod', 'CMS_recid', 'Catalogue_recid', 'Streamer_recid', 'Timestamp',
         'Conflicts', 'Albums::DDEX Status', 'Audio_Truth', 'Rights Territories', 'Track Number', 'Album Catalogue', 'Track Name',
         'Filename', 'Audio_Vision_URL', 'Audio_S3_URL', 'CMS_Audio_File', 'AudioHashSum'}
wb = load_workbook(xlsx, read_only=True)
rows = [r for r in wb['Needs review'].iter_rows(min_row=2, values_only=True) if r[5] and r[5].startswith('values disagree') and ' · ' not in r[5]]
plan, still = [], []
for r in rows:
    keep = B[str(r[3])]; drops = [B[x.strip()] for x in str(r[4]).split(',')]
    sets, real = {}, []
    for d in drops:
        for k, v in d.items():
            if k in LINKS or blank(v): continue
            cur = sets.get(k, keep.get(k))
            if blank(cur): sets[k] = v; continue
            if str(cur).strip() == str(v).strip(): continue
            if k in SAME and SAME[k](cur, v):
                b = better(k, cur, v)
                if str(b).strip() != str(keep.get(k, '')).strip(): sets[k] = b
            else: real.append(f'{k}: "{cur}" vs "{v}"')
    row = dict(cat=r[0], title=r[1], keep=str(keep['id']), drops=[str(d['id']) for d in drops], set=sets, expect={k: keep.get(k, '') for k in sets}, real=real)
    (still if real else plan).append(row)
json.dump(plan, open(f'{outdir}/plan.json', 'w'), indent=1, ensure_ascii=False)
json.dump(still, open(f'{outdir}/still-review.json', 'w'), indent=1, ensure_ascii=False)
print('candidates', len(rows), '| cosmetic only (batch 2):', len(plan), 'songs,', sum(len(p['drops']) for p in plan), 'extra records,',
      sum(1 for p in plan if p['set']), 'kept copies tidied | real differences (stay on review):', len(still))
print('tidied fields:', collections.Counter(k for p in plan for k in p['set']).most_common())
print('real differences by field:', collections.Counter(x.split(':')[0] for s in still for x in s['real']).most_common())
