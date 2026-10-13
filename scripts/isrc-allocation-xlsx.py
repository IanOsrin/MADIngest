"""scripts/isrc-allocation-xlsx.py — fill-in sheet for the ISRC allocator (Ian, 2026-10-13: "songs that are missing ISRC
but the rest of the album has ISRC … give this to the guy who allocates ISRC, in a way that he can fill columns" +
"Obviously where nothing exists in cache").

A song is listed when, in MadStreamer as it stands today (nightly copy + today's apply logs):
  - its album (catalogue + album title, as MAD groups albums) has at least one song WITH a valid ISRC, and
  - it has no valid ISRC itself, and
  - the metadata cache has no ISRC for it (same catalogue + title), and its MAM twin (catalogue + Filename) has none.
Songs whose title + artist the cache has on ANOTHER catalogue (with an ISRC) go to "Check first" — probably the same
recording, which must not get a second ISRC. Read-only.

   python3 scripts/isrc-allocation-xlsx.py tmp/isrc-allocation-tracks.json tmp/isrc-gap-compare.json tmp/metadata-extract.json tmp/isrc-gap-mam.json out.xlsx
"""
import json, sys, os, re, unicodedata, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Protection, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

tracks_f, cmp_f, cache_f, mam_f, out = sys.argv[1:6]
def V(v):
    s = re.sub(r'[\s-]', '', str(v or '')).upper()
    return s if re.fullmatch(r'[A-Z]{2}[A-Z0-9]{3}\d{7}', s) else ''
cat = lambda c: re.sub(r'[^A-Za-z0-9]', '', str(c or '')).upper()
asc = lambda s: unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
strict = lambda s: re.sub(r'[^a-z0-9]+', '', asc(s))
loose = lambda s: re.sub(r'[^a-z0-9]+', '', re.sub(r'\(.*?\)|\[.*?\]', ' ', asc(s)))

def log(folder):
    p = os.path.expanduser(f'~/Downloads/{folder}/apply-log.jsonl')
    return [json.loads(l) for l in open(p) if l.strip()] if os.path.exists(p) else []
restored = {r['id'] for r in json.load(open(os.path.expanduser('~/Downloads/madstreamer-cache-isrc/with-version.json')))}
isrc_now = {}
for f in ('madstreamer-isrc-fill', 'madstreamer-cache-isrc'):
    for j in log(f):
        if j.get('ok') and j.get('after') and not (f == 'madstreamer-cache-isrc' and j['recordId'] in restored): isrc_now[j['recordId']] = j['after']
title_now = {j['recordId']: j['after'] for j in log('madstreamer-title-fix') if j.get('ok') and j.get('after')}
mam_isrc_now = {}
for f in ('mam-isrc-fill', 'mam-cache-isrc'):
    for j in log(f):
        if j.get('ok') and j.get('after'): mam_isrc_now[j['recordId']] = j['after']
upc_removed = {r['recordId'] for r in json.load(open(os.path.expanduser('~/Downloads/madstreamer-upc-removal/backup-2026-10-13.json')))['records'] if str(r.get('upc') or '').strip()}

T = json.load(open(tracks_f))
for t in T:
    t['isrc'] = isrc_now.get(t['id']) or V(t['isrc'])
    t['title'] = title_now.get(t['id']) or t['title']
cmp_rows = json.load(open(cmp_f))['rows']
mam_of = {str(r['msId']): str(r['mamId']) for r in cmp_rows if r['where'] == 'MadStreamer' and r['mamId']}
mam_raw = {str(m['id']): m['isrcRaw'] for m in json.load(open(mam_f))}
mam_isrc = lambda mid: mam_isrc_now.get(mid) or V(mam_raw.get(mid))

cache = json.load(open(cache_f)); cache = cache if isinstance(cache, list) else cache['rows']
alb = collections.defaultdict(list); by_title = collections.defaultdict(list)
for c in cache:
    if c.get('catalogue') and c.get('track_name'):
        alb[cat(c['catalogue'])].append(c); by_title[loose(c['track_name'])].append(c)
def cache_isrc(t):
    for key in (strict, loose):
        hit = {V(c.get('isrc')) for c in alb.get(cat(t['cat']), []) if key(c['track_name']) and key(c['track_name']) == key(t['title'])} - {''}
        if hit: return hit
    return set()
def elsewhere(t):
    a = loose(t['artist'] or t['album_artist'])
    return [c for c in by_title.get(loose(t['title']), []) if cat(c['catalogue']) != cat(t['cat']) and V(c.get('isrc'))
            and a and (a in loose(c.get('track_artist')) or loose(c.get('track_artist')) in a or a in loose(c.get('album_artist')))]

albums = collections.defaultdict(list)
for t in T:
    if t['cat'] or t['title']: albums[(cat(t['cat']), (t['album'] or '').strip().lower())].append(t)

todo, check, partial = [], [], 0
for k, songs in albums.items():
    have = [s for s in songs if s['isrc']]; miss = [s for s in songs if not s['isrc'] and s['title']]
    if not have or not miss: continue
    partial += 1
    for s in miss:
        mid = mam_of.get(s['id'])
        if cache_isrc(s) or (mid and mam_isrc(mid)): continue          # an ISRC exists somewhere: not for allocation
        s['mam'] = mid; s['albumIsrcs'] = sorted(x['isrc'] for x in have); s['albumSize'] = len(songs); s['albumHave'] = len(have)
        s['hidden'] = s['id'] in upc_removed
        e = elsewhere(s)
        if e: s['else'] = '; '.join(dict.fromkeys(f"{c['catalogue']} — {c.get('album_title')} — {c.get('track_artist')} ({c['isrc']})" for c in e[:3])); check.append(s)
        else: todo.append(s)

key = lambda s: (cat(s['cat']), (s['album'] or '').lower(), int(re.sub(r'\D', '', str(s['seq'] or '')) or 0))
todo.sort(key=key); check.sort(key=key)

# ---- workbook ----
wb = Workbook()
HEAD = PatternFill('solid', fgColor='6D28D9'); FILL = PatternFill('solid', fgColor='FEF9C3'); FILLHEAD = PatternFill('solid', fgColor='CA8A04')
thin = Side(style='thin', color='D4D4D4'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
INFO = ['Catalogue', 'Album', 'Album artist', 'Barcode (UPC)', 'Label', 'Year', 'Track #', 'Song', 'Track artist', 'Version', 'Duration',
        'Composer(s)', 'Filename', 'Other ISRCs on this album', 'Songs on album / with ISRC', 'MadStreamer record', 'MAM record', 'Note']
FILLS = ['NEW ISRC', 'Allocated by', 'Date allocated', 'Allocator\'s notes']
IW = [13, 26, 20, 15, 14, 6, 6, 30, 20, 12, 9, 26, 13, 30, 10, 10, 9, 26]
FW = [16, 14, 13, 30]

def sheet(ws, rows, extra=None):
    cols = INFO + ([extra[0]] if extra else []) + FILLS
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        is_fill = c.value in FILLS
        c.font = Font(bold=True, color='FFFFFF'); c.fill = FILLHEAD if is_fill else HEAD; c.alignment = Alignment(wrap_text=True, vertical='top')
    widths = IW + ([extra[1]] if extra else []) + FW
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for s in rows:
        isr = s['albumIsrcs']; rng = isr[0] if len(isr) == 1 else f'{isr[0]} … {isr[-1]} ({len(isr)})'
        note = 'Album hidden on the website until these ISRCs are in (barcodes removed 13 Oct)' if s['hidden'] else ''
        line = [s['cat'], s['album'], s['album_artist'], s['upc'], s['label'], s['yr'], s['seq'], s['title'], s['artist'], s['version'], s['duration'],
                s['composer'], s['file'], rng, f"{s['albumSize']} / {s['albumHave']}", int(s['id']), int(s['mam']) if s.get('mam') else '', note]
        if extra: line.append(s['else'])
        ws.append(line + ['', '', '', ''])
    n = ws.max_row; first_fill = len(cols) - len(FILLS) + 1
    L = get_column_letter(first_fill)
    for r in range(2, n + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(r, c); cell.border = box
            if c >= first_fill: cell.fill = FILL; cell.protection = Protection(locked=False)
    # ISRC must be 12 characters CC-XXX-YY-NNNNN (dashes/spaces allowed)
    dv = DataValidation(type='custom', allow_blank=True, showErrorMessage=True, errorTitle='Not an ISRC',
                        error='An ISRC is 12 characters: 2 letters, 3 letters/digits, 7 digits — e.g. ZAC032600123 (dashes allowed).',
                        formula1=f'LEN(SUBSTITUTE(SUBSTITUTE({L}2,"-","")," ",""))=12')
    dv.add(f'{L}2:{L}{max(n, 2)}'); ws.add_data_validation(dv)
    dd = DataValidation(type='date', operator='greaterThan', formula1='45000', allow_blank=True, showErrorMessage=True, error='Please enter a date')
    D = get_column_letter(first_fill + 2); dd.add(f'{D}2:{D}{max(n, 2)}'); ws.add_data_validation(dd)
    for r in range(2, n + 1): ws.cell(r, first_fill + 2).number_format = 'yyyy-mm-dd'
    # flag an ISRC typed twice in the sheet
    ws.conditional_formatting.add(f'{L}2:{L}{max(n, 2)}', FormulaRule(formula=[f'AND({L}2<>"",COUNTIF(${L}$2:${L}${max(n, 2)},{L}2)>1)'], fill=PatternFill('solid', fgColor='FCA5A5')))
    ws.freeze_panes = 'I2'; ws.auto_filter.ref = ws.dimensions
    ws.protection.sheet = True; ws.protection.autoFilter = False; ws.protection.sort = False; ws.protection.formatColumns = False
    ws.protection.formatRows = False                     # no password: Review ▸ Unprotect Sheet if needed

ws = wb.active; ws.title = 'To allocate'; sheet(ws, todo)
ws2 = wb.create_sheet('Check first'); sheet(ws2, check, ('The cache has this title + artist on ANOTHER catalogue', 50))

al = wb.create_sheet('Albums')
al.append(['Catalogue', 'Album', 'Album artist', 'Barcode (UPC)', 'Songs on album', 'Have ISRC', 'To allocate', 'Check first', 'Hidden on website?'])
for c in al[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD
for i, w in enumerate([13, 34, 24, 15, 9, 9, 10, 10, 12], 1): al.column_dimensions[get_column_letter(i)].width = w
g = collections.OrderedDict()
for s in todo + check: g.setdefault((cat(s['cat']), (s['album'] or '').lower()), []).append(s)
for k, ss in sorted(g.items()):
    s = ss[0]
    al.append([s['cat'], s['album'], s['album_artist'], s['upc'], s['albumSize'], s['albumHave'], sum(1 for x in ss if x in todo), sum(1 for x in ss if x in check),
               'yes — barcodes removed 13 Oct' if any(x['hidden'] for x in ss) else 'no'])
al.freeze_panes = 'C2'; al.auto_filter.ref = al.dimensions

rd = wb.create_sheet('How to fill')
for l in [
    'Songs that need an ISRC allocated: each is on an album (catalogue + album title) where the other songs already have ISRCs, and no ISRC exists for it anywhere — not in MadStreamer (the website\'s database), not in MAM, and not in the metadata cache (Ingrooves extract + ISRC register). Taken 13 Oct 2026.',
    'Fill only the YELLOW columns: NEW ISRC, Allocated by, Date allocated (yyyy-mm-dd), Allocator\'s notes. The other columns are locked so they can\'t be changed by accident (Review ▸ Unprotect Sheet — no password — if you need to). Sorting and filtering work.',
    'NEW ISRC accepts 12 characters (dashes or spaces allowed), e.g. ZAC032600123 or ZA-C03-26-00123. The same ISRC typed twice in the sheet turns red.',
    '"Other ISRCs on this album" shows the album\'s existing ISRCs (first … last), to help keep the album\'s numbering together.',
    '"Check first": the cache has the same song title by the same artist on ANOTHER catalogue, with an ISRC. That is probably the same recording, which must NOT get a second ISRC — if it is the same recording, write that ISRC in NEW ISRC and say so in the notes; if it is a different recording (re-recording, live, edit), allocate a new one.',
    'Note column: "Album hidden on the website" = we removed this album\'s barcode on 13 Oct because some songs lacked ISRCs; it goes back on the website once these are filled in.',
    'Send the file back as it is — the ISRCs will be loaded into MadStreamer (and MAM, using the record numbers) from it.',
]: rd.append([l])
rd.column_dimensions['A'].width = 140
for c in rd['A']: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.move_sheet('How to fill', offset=-3)
wb.active = 1
wb.save(out)
print(out, '| partial albums', partial, '| to allocate', len(todo), 'on', len({key(s)[:2] for s in todo}), 'albums | check first', len(check), '| hidden', sum(1 for s in todo + check if s['hidden']))
