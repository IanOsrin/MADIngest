"""scripts/isrc-cache-truth.py — treat the metadata cache as the ONLY truth for ISRCs and list what MadStreamer and
MAM would have to change to match it (Ian, 2026-10-13: "accept that only the cache is the source of truth").
Read-only.

Every MadStreamer / MAM song is matched to the cache by catalogue + title: exact title first (ignoring case,
accents, punctuation), then ignoring bracketed bits, each only when it picks ONE ISRC on that cache album. Today's
fills (madstreamer-isrc-fill, mam-isrc-fill apply logs) are laid over the scanned values.

   python3 scripts/isrc-cache-truth.py tmp/isrc-gap-compare.json tmp/metadata-extract.json tmp/isrc-gap-mam.json out.xlsx
"""
import json, sys, os, re, unicodedata, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

cmp_, cache_f, mam_f, out = sys.argv[1:5]
def V(v):
    s = re.sub(r'[\s-]', '', str(v or '')).upper()
    return s if re.fullmatch(r'[A-Z]{2}[A-Z0-9]{3}\d{7}', s) else ''
cat = lambda c: re.sub(r'[^A-Za-z0-9]', '', str(c or '')).upper()
asc = lambda s: unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
strict = lambda s: re.sub(r'[^a-z0-9]+', '', asc(s))
loose = lambda s: re.sub(r'[^a-z0-9]+', '', re.sub(r'\(.*?\)|\[.*?\]', ' ', asc(s)))

rows = [r for r in json.load(open(cmp_))['rows'] if r['cat'] or r['title']]
cache = json.load(open(cache_f)); cache = cache if isinstance(cache, list) else cache['rows']
mam_all = json.load(open(mam_f))

# today's fills, laid over the scan
def filled(folder, idkey):
    p = os.path.expanduser(f'~/Downloads/{folder}/apply-log.jsonl')
    return {json.loads(l)['recordId']: json.loads(l)['after'] for l in open(p) if l.strip() and json.loads(l).get('ok') and json.loads(l).get('after')} if os.path.exists(p) else {}
ms_fill, mam_fill = filled('madstreamer-isrc-fill', 'fm_record_id'), filled('mam-isrc-fill', 'recordId')

# the cache, by album
alb = collections.defaultdict(list)
for c in cache:
    if c.get('catalogue') and c.get('track_name'): alb[cat(c['catalogue'])].append(c)
reg_titles = collections.defaultdict(set)                    # ISRC → distinct song titles in the cache
for c in cache:
    if V(c.get('isrc')): reg_titles[V(c['isrc'])].add(loose(c['track_name']) or strict(c['track_name']))

def truth(c_, title):
    """→ (status, cache row) ; status: match | ambiguous | song not in cache album | album not in cache"""
    a = alb.get(cat(c_))
    if not a: return 'album not in cache', None
    for key in (strict, loose):
        hit = [c for c in a if key(c['track_name']) and key(c['track_name']) == key(title)]
        isrcs = {V(c.get('isrc')) for c in hit}
        if len(isrcs) == 1: return 'match', hit[0]
        if len(isrcs) > 1: return 'ambiguous', hit[0]
    return 'song not in cache album', None

# current values
ms = []
for r in rows:
    if r['where'] != 'MadStreamer': continue
    v = ms_fill.get(str(r['msId'])) or r['ms']
    ms.append(dict(db='MadStreamer', id=r['msId'], cat=r['cat'], album=r['album'], title=r['title'], artist=r['artist'], seq=r['seq'], file=r['file'],
                   cur=v or '', raw=v or (r['msRaw'] or ''), onSite=r['onSite']))
mam = [dict(db='MAM', id=m['id'], cat=m['cat'], album='', title=m['title'], artist='', seq=m['seq'], file=m['file'],
            cur=mam_fill.get(str(m['id'])) or V(m['isrcRaw']), raw=mam_fill.get(str(m['id'])) or m['isrcRaw'], onSite=False) for m in mam_all]

def held_by(songs):                                          # (catalogue, ISRC) → titles carrying it in that db
    h = collections.defaultdict(list)
    for s in songs:
        if s['cur']: h[(cat(s['cat']), s['cur'])].append(s['title'] or '(no title)')
    return h

results = {}
for name, songs in (('MadStreamer', ms), ('MAM', mam)):
    hb = held_by(songs); res = []
    for s in songs:
        st, c = truth(s['cat'], s['title'])
        s.update(status=st, c=c, taken='', regMany=False, truth='')
        if st == 'match':
            t = V(c.get('isrc'))
            if not t: s['verdict'] = 'cache has no ISRC either'
            elif s['cur'] == t: s['verdict'] = 'agrees'
            elif not s['cur']: s['verdict'] = 'fill from cache'
            else: s['verdict'] = 'replace with cache'
            if t:
                s['truth'] = t
                others = [x for x in hb.get((cat(s['cat']), t), []) if strict(x) != strict(s['title'])]
                s['taken'] = '; '.join(dict.fromkeys(others))
                s['regMany'] = len(reg_titles[t]) > 1
        else: s['verdict'] = st
        res.append(s)
    results[name] = res

# ---- workbook ----
wb = Workbook()
H = PatternFill('solid', fgColor='6D28D9')
def head(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = H; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'D2'

s = wb.active; s.title = 'Summary'
s.append(['ISRCs with the metadata cache as the only truth (13 Oct 2026)']); s['A1'].font = Font(bold=True, size=13)
order = ['agrees', 'fill from cache', 'replace with cache', 'cache has no ISRC either', 'ambiguous', 'song not in cache album', 'album not in cache']
label = {'agrees': 'Already agrees with the cache', 'fill from cache': 'Has NO ISRC — the cache has one (fill)',
         'replace with cache': 'Has a DIFFERENT ISRC from the cache (replace)', 'cache has no ISRC either': 'Matched, but the cache has no ISRC either',
         'ambiguous': 'Cache album has this title twice with different ISRCs (can\'t tell which)', 'song not in cache album': 'Album is in the cache but this song isn\'t',
         'album not in cache': 'Album not in the cache at all'}
s.append([]); s.append(['', 'MadStreamer', 'MAM'])
for k in order: s.append([label[k], sum(1 for x in results['MadStreamer'] if x['verdict'] == k), sum(1 for x in results['MAM'] if x['verdict'] == k)])
s.append(['Songs in total', len(results['MadStreamer']), len(results['MAM'])])
s.append([])
ch = lambda db: [x for x in results[db] if x['verdict'] in ('fill from cache', 'replace with cache')]
s.append(['Changes needed to match the cache', len(ch('MadStreamer')), len(ch('MAM'))])
s.append(['  …on songs the website shows now (ISRC + barcode, not hidden)', sum(1 for x in ch('MadStreamer') if x['onSite'] and x['cur']), ''])
s.append(['  …where the cache ISRC is currently on ANOTHER song of that album in that database (fixes come in pairs/chains — do the whole album)',
          sum(1 for x in ch('MadStreamer') if x['taken']), sum(1 for x in ch('MAM') if x['taken'])])
s.append(['  …where the cache itself has that ISRC on more than one different song (cache may be wrong — check)',
          sum(1 for x in ch('MadStreamer') if x['regMany']), sum(1 for x in ch('MAM') if x['regMany'])])
s.append(['  …albums affected', len({cat(x['cat']) for x in ch('MadStreamer')}), len({cat(x['cat']) for x in ch('MAM')})])
s.column_dimensions['A'].width = 100; s.column_dimensions['B'].width = 13; s.column_dimensions['C'].width = 13

COLS = ['Catalogue', 'Album', 'Song', 'Artist', '#', 'Filename', 'Record', 'Change', 'ISRC now', 'Cache ISRC (truth)', 'Cache title', 'Cache #',
        'Cache ISRC now on another song here', 'Cache has this ISRC on several songs?', 'On website now?', 'Done ✓']
W = [14, 26, 28, 20, 4, 13, 9, 10, 15, 16, 28, 6, 34, 12, 9, 7]
for db in ('MadStreamer', 'MAM'):
    ws = wb.create_sheet(f'{db} changes'); head(ws, COLS, W)
    for x in sorted(ch(db), key=lambda x: (cat(x['cat']), str(x['seq']).zfill(4))):
        c = x['c']
        ws.append([x['cat'], x['album'] or c.get('album_title'), x['title'], x['artist'] or c.get('track_artist'), x['seq'], x['file'], int(x['id']),
                   'fill' if x['verdict'] == 'fill from cache' else 'replace', x['raw'] or '—', x['truth'], c['track_name'], c.get('seq'),
                   x['taken'], 'YES — check' if x['regMany'] else '', ('yes' if x['onSite'] and x['cur'] else 'no') if db == 'MadStreamer' else '', ''])
    ws.auto_filter.ref = ws.dimensions

# the cache's own doubtful entries that touch these changes
ws = wb.create_sheet('Cache doubts')
head(ws, ['Database', 'Catalogue', 'Song', 'Record', 'Why it was not decided', 'Cache rows for this title'], [12, 14, 32, 9, 46, 80])
for db in ('MadStreamer', 'MAM'):
    for x in results[db]:
        if x['verdict'] == 'ambiguous':
            a = [c for c in alb[cat(x['cat'])] if strict(c['track_name']) == strict(x['title']) or loose(c['track_name']) == loose(x['title'])]
            ws.append([db, x['cat'], x['title'], int(x['id']), 'cache album has this title more than once with different ISRCs',
                       '; '.join(f"{c['track_name']} #{c.get('seq')} = {c.get('isrc')}" for c in a)])
ws.auto_filter.ref = ws.dimensions

# MadStreamer songs the cache can't speak for, by album
ws = wb.create_sheet('MadStreamer not in cache')
head(ws, ['Catalogue', 'Album', 'Songs', 'Why', 'Of those with an ISRC now', 'On website now'], [15, 40, 8, 30, 12, 12])
g = collections.defaultdict(list)
for x in results['MadStreamer']:
    if x['verdict'] in ('album not in cache', 'song not in cache album'): g[(x['cat'], x['album'], x['verdict'])].append(x)
for (c_, a, v), xs in sorted(g.items(), key=lambda kv: -len(kv[1])):
    ws.append([c_, a, len(xs), v, sum(1 for x in xs if x['cur']), sum(1 for x in xs if x['onSite'] and x['cur'])])
ws.auto_filter.ref = ws.dimensions

rd = wb.create_sheet('Read me')
for l in [
    'Rule: the metadata cache (Ingrooves extract + ISRC register, 82,540 rows) is the ONLY truth. Every song in MadStreamer (website copy of ~03:00 today, plus today\'s 690 fills) and MAM (read live today, plus today\'s 128 fills) is compared with the cache\'s ISRC for the same song. Nothing was changed.',
    'Matching: same catalogue (ignoring spaces/punctuation) + same song title — exactly first (ignoring case, accents, punctuation), then ignoring bracketed bits like (Remix), each only when the title gives ONE ISRC on that cache album. Titles the cache album has twice with different ISRCs are on "Cache doubts" and are not decided.',
    '"fill" = the database has no valid ISRC and the cache has one. "replace" = the database has a different ISRC.',
    '"Cache ISRC now on another song here": in that database another song of the same catalogue already carries the cache\'s ISRC — usually an album whose ISRCs slid a track. Those must be changed album by album, together, or the album briefly holds the same ISRC twice.',
    '"Cache has this ISRC on several songs?": the cache itself registers that ISRC to more than one different title, so the cache may be wrong there — check before changing.',
    'Songs whose album or title the cache doesn\'t have are counted on the Summary; MadStreamer ones are listed by album on "MadStreamer not in cache".',
]: rd.append([l])
rd.column_dimensions['A'].width = 150
for c in rd['A']: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(out)
for db in ('MadStreamer', 'MAM'):
    print(db, collections.Counter(x['verdict'] for x in results[db]).most_common())
print(out)
